import os
import re
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas

BASE_DIR = Path(r"C:\Users\chakr\.gemini\antigravity\scratch\Need2Done\docs\NEED2DONE_LAUNCH")

# -------------------------------------------------------------
# Color Palette
# -------------------------------------------------------------
NAVY_HEX = "022B5E"
ORANGE_HEX = "FF7D00"
DARK_GRAY_HEX = "1E293B"
LIGHT_BG_HEX = "F1F5F9"
BORDER_HEX = "CBD5E1"

NAVY_COLOR = colors.HexColor(f"#{NAVY_HEX}")
ORANGE_COLOR = colors.HexColor(f"#{ORANGE_HEX}")
DARK_TEXT = colors.HexColor(f"#{DARK_GRAY_HEX}")
LIGHT_BG = colors.HexColor(f"#{LIGHT_BG_HEX}")
BORDER_COLOR = colors.HexColor(f"#{BORDER_HEX}")


# -------------------------------------------------------------
# Markdown Parser Helper
# -------------------------------------------------------------
def parse_markdown(md_text):
    lines = md_text.splitlines()
    blocks = []
    i = 0
    in_code = False
    code_lines = []

    while i < len(lines):
        line = lines[i]

        # Code Block
        if line.strip().startswith("```"):
            if in_code:
                blocks.append(('code', "\n".join(code_lines)))
                code_lines = []
                in_code = False
            else:
                in_code = True
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        stripped = line.strip()

        # Empty line
        if not stripped:
            i += 1
            continue

        # Horizontal Rule
        if re.match(r'^-{3,}$', stripped) or re.match(r'^\*{3,}$', stripped):
            blocks.append(('hr', ''))
            i += 1
            continue

        # Table detection
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Parse table rows
            rows = []
            for tl in table_lines:
                # Check if separator row
                if re.match(r'^\|[\s\-:|]+\|$', tl):
                    continue
                cells = [c.strip() for c in tl[1:-1].split('|')]
                rows.append(cells)
            if rows:
                blocks.append(('table', rows))
            continue

        # Headings
        if stripped.startswith('# '):
            blocks.append(('h1', stripped[2:].strip()))
            i += 1
            continue
        if stripped.startswith('## '):
            blocks.append(('h2', stripped[3:].strip()))
            i += 1
            continue
        if stripped.startswith('### '):
            blocks.append(('h3', stripped[4:].strip()))
            i += 1
            continue
        if stripped.startswith('#### '):
            blocks.append(('h4', stripped[5:].strip()))
            i += 1
            continue

        # Blockquote
        if stripped.startswith('> '):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                quote_lines.append(lines[i].strip()[1:].strip())
                i += 1
            blocks.append(('quote', " ".join(quote_lines)))
            continue

        # Checklist
        if re.match(r'^-\s+\[([ xX])\]\s+(.*)', stripped):
            m = re.match(r'^-\s+\[([ xX])\]\s+(.*)', stripped)
            checked = m.group(1).lower() == 'x'
            blocks.append(('checklist', (checked, m.group(2))))
            i += 1
            continue

        # Bullet List
        if re.match(r'^[\-\*\+]\s+(.*)', stripped):
            m = re.match(r'^[\-\*\+]\s+(.*)', stripped)
            blocks.append(('bullet', m.group(1)))
            i += 1
            continue

        # Numbered List
        if re.match(r'^\d+\.\s+(.*)', stripped):
            m = re.match(r'^\d+\.\s+(.*)', stripped)
            blocks.append(('numbered', m.group(1)))
            i += 1
            continue

        # Paragraph (may span multiple non-empty lines)
        para_lines = [stripped]
        i += 1
        while i < len(lines):
            next_line = lines[i].strip()
            if not next_line:
                break
            if next_line.startswith('#') or next_line.startswith('|') or next_line.startswith('```') or next_line.startswith('>') or re.match(r'^[\-\*\+]\s+', next_line) or re.match(r'^\d+\.\s+', next_line):
                break
            para_lines.append(next_line)
            i += 1
        blocks.append(('p', " ".join(para_lines)))

    return blocks


def clean_markdown_inline(text):
    # Remove file links [text](url) -> text
    text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)
    return text


# -------------------------------------------------------------
# DOCX Generator
# -------------------------------------------------------------
def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''<w:tcMar {nsdecls("w")}>
        <w:top w:w="{top}" w:type="dxa"/>
        <w:bottom w:w="{bottom}" w:type="dxa"/>
        <w:left w:w="{left}" w:type="dxa"/>
        <w:right w:w="{right}" w:type="dxa"/>
    </w:tcMar>''')
    tcPr.append(tcMar)

def add_styled_inline_runs(paragraph, text, default_size=10.5, default_color=(30, 41, 59)):
    text = clean_markdown_inline(text)
    # Split by bold **
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(default_size - 1)
        else:
            run = paragraph.add_run(part)
        
        run.font.size = Pt(default_size)
        run.font.color.rgb = RGBColor(*default_color)

def generate_docx(blocks, output_path, title="Need2Done Documentation"):
    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Need2Done | Commercial Launch & Compliance Document")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Confidential & Proprietary • Need2Done (N2D) India")
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

    # Populate Blocks
    for b_type, b_content in blocks:
        if b_type == 'h1':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(clean_markdown_inline(b_content))
            run.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(2, 43, 94) # Navy

        elif b_type == 'h2':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(clean_markdown_inline(b_content))
            run.bold = True
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(2, 43, 94)

        elif b_type == 'h3':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(clean_markdown_inline(b_content))
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(255, 125, 0) # Orange

        elif b_type == 'h4':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(clean_markdown_inline(b_content))
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(71, 85, 105)

        elif b_type == 'p':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            add_styled_inline_runs(p, b_content)

        elif b_type == 'bullet':
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_styled_inline_runs(p, b_content)

        elif b_type == 'numbered':
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_styled_inline_runs(p, b_content)

        elif b_type == 'checklist':
            checked, text = b_content
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(3)
            mark_run = p.add_run("[✓] " if checked else "[  ] ")
            mark_run.bold = True
            mark_run.font.color.rgb = RGBColor(16, 185, 129) if checked else RGBColor(100, 116, 139)
            add_styled_inline_runs(p, text)

        elif b_type == 'quote':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("┃  ")
            run.bold = True
            run.font.color.rgb = RGBColor(255, 125, 0)
            add_styled_inline_runs(p, b_content, default_color=(71, 85, 105))

        elif b_type == 'code':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(b_content)
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(15, 23, 42)

        elif b_type == 'hr':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("—" * 55)
            run.font.color.rgb = RGBColor(226, 232, 240)

        elif b_type == 'table':
            rows_data = b_content
            if not rows_data:
                continue
            table = doc.add_table(rows=len(rows_data), cols=len(rows_data[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True

            for r_idx, row in enumerate(rows_data):
                for c_idx, cell_value in enumerate(row):
                    if c_idx >= len(table.rows[r_idx].cells):
                        continue
                    cell = table.cell(r_idx, c_idx)
                    cell.text = ""
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_after = Pt(2)
                    cp.paragraph_format.space_before = Pt(2)

                    if r_idx == 0:
                        set_cell_background(cell, NAVY_HEX)
                        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
                        add_styled_inline_runs(cp, cell_value, default_size=9.5, default_color=(255, 255, 255))
                        cp.runs[0].bold = True if cp.runs else None
                    else:
                        bg_hex = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                        set_cell_background(cell, bg_hex)
                        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                        add_styled_inline_runs(cp, cell_value, default_size=9)

            doc.add_paragraph() # spacing after table

    doc.save(str(output_path))
    print(f"[DOCX OK] {output_path.name}")


# -------------------------------------------------------------
# PDF Generator using ReportLab
# -------------------------------------------------------------
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super(NumberedCanvas, self).showPage()
        super(NumberedCanvas, self).save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        
        # Header (pages after first)
        if self._pageNumber > 1:
            self.drawString(54, letter[1] - 36, "Need2Done — Commercial Launch & Legal Compliance")
            self.setStrokeColor(colors.HexColor("#E2E8F0"))
            self.setLineWidth(0.5)
            self.line(54, letter[1] - 42, letter[0] - 54, letter[1] - 42)

        # Footer
        self.drawString(54, 36, "Confidential & Proprietary • Need2Done (N2D) India")
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(letter[0] - 54, 36, page_str)
        self.setStrokeColor(colors.HexColor("#E2E8F0"))
        self.setLineWidth(0.5)
        self.line(54, 46, letter[0] - 54, 46)

        self.restoreState()


def format_inline_html(text):
    text = clean_markdown_inline(text)
    # Convert bold **text** to <b>text</b>
    text = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', text)
    # Convert italic *text* to <i>text</i>
    text = re.sub(r'\*(.*?)\*', r'<i>\1</i>', text)
    # Convert code `code` to <font name="Courier">code</font>
    text = re.sub(r'`(.*?)`', r'<font name="Courier" size="8.5" color="#0F172A">\1</font>', text)
    # Replace & that aren't entities
    text = re.sub(r'&(?!(amp|lt|gt|quot|apos);)', '&amp;', text)
    return text


def generate_pdf(blocks, output_path, title="Need2Done Documentation"):
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=NAVY_COLOR,
        spaceBefore=10,
        spaceAfter=8
    )

    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=NAVY_COLOR,
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True
    )

    h3_style = ParagraphStyle(
        'DocH3',
        parent=styles['Heading3'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=ORANGE_COLOR,
        spaceBefore=10,
        spaceAfter=4,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=DARK_TEXT,
        spaceAfter=6
    )

    bullet_style = ParagraphStyle(
        'DocBullet',
        parent=body_style,
        leftIndent=16,
        firstLineIndent=-10,
        spaceAfter=3
    )

    quote_style = ParagraphStyle(
        'DocQuote',
        parent=body_style,
        fontName='Helvetica-Oblique',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#475569'),
        leftIndent=18,
        spaceBefore=4,
        spaceAfter=6
    )

    code_style = ParagraphStyle(
        'DocCode',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=8,
        leading=11,
        textColor=colors.HexColor('#0F172A'),
        backColor=LIGHT_BG,
        borderPadding=6,
        spaceBefore=4,
        spaceAfter=6
    )

    table_header_style = ParagraphStyle(
        'TableHeader',
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
        alignment=0
    )

    table_cell_style = ParagraphStyle(
        'TableCell',
        fontName='Helvetica',
        fontSize=8,
        leading=10.5,
        textColor=DARK_TEXT,
        alignment=0
    )

    story = []

    for b_type, b_content in blocks:
        if b_type == 'h1':
            story.append(Paragraph(format_inline_html(b_content), title_style))
        elif b_type == 'h2':
            story.append(Paragraph(format_inline_html(b_content), h2_style))
        elif b_type == 'h3':
            story.append(Paragraph(format_inline_html(b_content), h3_style))
        elif b_type == 'h4':
            story.append(Paragraph(format_inline_html(b_content), h3_style))
        elif b_type == 'p':
            story.append(Paragraph(format_inline_html(b_content), body_style))
        elif b_type == 'bullet':
            story.append(Paragraph(f"• {format_inline_html(b_content)}", bullet_style))
        elif b_type == 'numbered':
            story.append(Paragraph(f"1. {format_inline_html(b_content)}", bullet_style))
        elif b_type == 'checklist':
            checked, text = b_content
            mark = '<font color="#10B981"><b>[✓]</b></font>' if checked else '<font color="#64748B"><b>[ ]</b></font>'
            story.append(Paragraph(f"{mark} {format_inline_html(text)}", bullet_style))
        elif b_type == 'quote':
            story.append(Paragraph(f"<b>|</b> {format_inline_html(b_content)}", quote_style))
        elif b_type == 'code':
            # Escape HTML
            escaped_code = b_content.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('\n', '<br/>')
            story.append(Paragraph(escaped_code, code_style))
        elif b_type == 'hr':
            story.append(HRFlowable(width="100%", thickness=0.8, color=BORDER_COLOR, spaceBefore=6, spaceAfter=8))
        elif b_type == 'table':
            rows_data = b_content
            if not rows_data:
                continue

            num_cols = len(rows_data[0])
            total_w = letter[0] - 108 # 504 pt
            col_w = total_w / num_cols

            flowable_table_data = []
            for r_idx, row in enumerate(rows_data):
                row_cells = []
                for cell_text in row:
                    style_to_use = table_header_style if r_idx == 0 else table_cell_style
                    p_cell = Paragraph(format_inline_html(cell_text), style_to_use)
                    row_cells.append(p_cell)
                flowable_table_data.append(row_cells)

            t = Table(flowable_table_data, colWidths=[col_w] * num_cols)
            t_style = [
                ('BACKGROUND', (0, 0), (-1, 0), NAVY_COLOR),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                ('LEFTPADDING', (0, 0), (-1, -1), 5),
                ('RIGHTPADDING', (0, 0), (-1, -1), 5),
                ('GRID', (0, 0), (-1, -1), 0.5, BORDER_COLOR),
            ]
            for r_idx in range(1, len(rows_data)):
                if r_idx % 2 == 1:
                    t_style.append(('BACKGROUND', (0, r_idx), (-1, r_idx), LIGHT_BG))
            t.setStyle(TableStyle(t_style))
            story.append(Spacer(1, 4))
            story.append(t)
            story.append(Spacer(1, 8))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"[PDF OK]  {output_path.name}")


# -------------------------------------------------------------
# Main Conversion Engine
# -------------------------------------------------------------
def main():
    print(f"Scanning directory: {BASE_DIR}")
    md_files = list(BASE_DIR.glob("**/*.md"))
    print(f"Found {len(md_files)} markdown files.")

    all_master_blocks = [
        ('h1', 'Need2Done (N2D) Commercial Launch Master Dossier'),
        ('h3', 'Complete Legal Policies, Terms of Service, Partner Agreements & Statutory Checklists'),
        ('p', '<b>Version:</b> 3.0 Production Ready • <b>Date:</b> October 2026 • <b>Jurisdiction:</b> India'),
        ('hr', '')
    ]

    for md_file in sorted(md_files):
        print(f"\nProcessing: {md_file.relative_to(BASE_DIR)}")
        with open(md_file, 'r', encoding='utf-8') as f:
            content = f.read()

        blocks = parse_markdown(content)

        docx_path = md_file.with_suffix('.docx')
        pdf_path = md_file.with_suffix('.pdf')

        generate_docx(blocks, docx_path, title=md_file.stem)
        generate_pdf(blocks, pdf_path, title=md_file.stem)

        if md_file.name != "README.md":
            all_master_blocks.extend(blocks)
            all_master_blocks.append(('hr', ''))

    # Generate Master Unified Dossier
    master_docx = BASE_DIR / "Need2Done_Commercial_Launch_Master_Dossier.docx"
    master_pdf = BASE_DIR / "Need2Done_Commercial_Launch_Master_Dossier.pdf"
    
    print("\nGenerating Master Combined Dossier...")
    generate_docx(all_master_blocks, master_docx, title="Need2Done Commercial Launch Master Dossier")
    generate_pdf(all_master_blocks, master_pdf, title="Need2Done Commercial Launch Master Dossier")

    print("\nAll conversions completed successfully!")

if __name__ == "__main__":
    main()
